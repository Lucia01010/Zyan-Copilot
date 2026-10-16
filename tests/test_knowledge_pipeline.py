"""M17 Knowledge Pipeline + M18 Retrieval Intelligence 测试。

M17.1 多源 Loader（Markdown/Excel/Database/API）、M17.2 Document metadata、
M18.1 Query Rewrite、M18.2 Hybrid Retrieval、M18.3 Rerank（粗排→精排）。
"""
import sqlite3
import tempfile
import unittest
from pathlib import Path

from app.knowledge_base import KnowledgeStore
from app.knowledge_loader import Document, KnowledgeLoader, parse_front_matter
from app.query_rewriter import rewrite_query


class FrontMatterTests(unittest.TestCase):
    def test_front_matter_parsed_and_stripped(self):
        meta, body = parse_front_matter("---\ncategory: 产品知识\nconfidence: 0.95\n---\n# 标题\n正文")
        self.assertEqual(meta["category"], "产品知识")
        self.assertEqual(meta["confidence"], 0.95)
        self.assertTrue(body.startswith("# 标题"))

    def test_no_front_matter_returns_original(self):
        meta, body = parse_front_matter("# 标题\n正文")
        self.assertEqual(meta, {})
        self.assertTrue(body.startswith("# 标题"))


class LoaderTests(unittest.TestCase):
    def setUp(self):
        self.loader = KnowledgeLoader()
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_markdown_loader_carries_front_matter_metadata(self):
        path = self.dir / "guide.md"
        path.write_text("---\ncategory: 产品知识\nsource: 产品说明书\n---\n# 护肤指南\n\n## 精华使用\n早晚洁面后取 2-3 滴。", encoding="utf-8")
        documents = self.loader.load_markdown(path)
        self.assertTrue(documents)
        for doc in documents:
            self.assertEqual(doc.metadata["category"], "产品知识")
            self.assertEqual(doc.metadata["source"], "产品说明书")
            self.assertEqual(doc.metadata["doc_title"], "护肤指南")
        section = next(doc for doc in documents if doc.metadata["section"] == "精华使用")
        self.assertIn("2-3 滴", section.text)

    def test_excel_loader_rows_become_documents(self):
        from openpyxl import Workbook

        path = self.dir / "成分表.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["text", "category", "product_line"])
        sheet.append(["修护精华液含 B5 泛醇，帮助屏障修护。", "产品知识", "修护系列"])
        sheet.append(["防晒霜需出门前 15 分钟涂抹。", "产品知识", "防晒系列"])
        workbook.save(path)
        workbook.close()
        documents = self.loader.load_excel(path)
        self.assertEqual(len(documents), 2)
        self.assertIn("B5 泛醇", documents[0].text)
        self.assertEqual(documents[0].metadata["category"], "产品知识")
        self.assertEqual(documents[0].metadata["product_line"], "修护系列")
        self.assertEqual(documents[1].metadata["product_line"], "防晒系列")

    def test_excel_loader_skips_rows_without_text(self):
        from openpyxl import Workbook

        path = self.dir / "activity.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["text", "category"])
        sheet.append(["满 300 减 50。", "活动规则"])
        sheet.append(["", "活动规则"])
        workbook.save(path)
        workbook.close()
        documents = self.loader.load_excel(path)
        self.assertEqual(len(documents), 1)

    def test_database_loader_reads_only_knowledge_table(self):
        db_path = self.dir / "source.sqlite3"
        conn = sqlite3.connect(db_path)
        conn.execute("CREATE TABLE policy_rules (rule_text TEXT, effective_date TEXT)")
        conn.execute("INSERT INTO policy_rules VALUES ('未拆封商品支持七天无理由退货。', '2026-01-01')")
        conn.execute("INSERT INTO policy_rules VALUES (NULL, '2026-01-02')")
        conn.commit()
        conn.close()
        documents = self.loader.load_database(db_path, "policy_rules", "rule_text", ("effective_date",), category="售后规则")
        self.assertEqual(len(documents), 1)
        self.assertIn("七天无理由", documents[0].text)
        self.assertEqual(documents[0].metadata["effective_date"], "2026-01-01")
        self.assertEqual(documents[0].metadata["source"], "database:policy_rules")

    def test_api_loader_injected_fetch_and_no_network_default(self):
        documents = self.loader.load_api("activity_rules", fetch=lambda: [{"text": "618 预售定金可抵 100 元。", "metadata": {"category": "活动规则"}}])
        self.assertEqual(len(documents), 1)
        self.assertEqual(documents[0].metadata["source"], "api:activity_rules")
        self.assertEqual(self.loader.load_api("offline"), [])

    def test_unsupported_extension_rejected(self):
        with self.assertRaises(ValueError):
            self.loader.load_path(self.dir / "unknown.txt")


class QueryRewriteTests(unittest.TestCase):
    def test_colloquial_pain_rewritten_to_standard_terms(self):
        result = rewrite_query("这个用了脸疼怎么办")
        self.assertIn("刺痛", result["added_terms"])
        self.assertIn("过敏", result["added_terms"])
        self.assertIn("刺痛", result["rewritten"])
        self.assertEqual(result["original"], "这个用了脸疼怎么办")

    def test_no_match_keeps_original(self):
        result = rewrite_query("物流到哪了")
        self.assertEqual(result["added_terms"], [])
        self.assertEqual(result["rewritten"], "物流到哪了")


class RetrievalIntelligenceTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)
        self.kb_dir = self.dir / "knowledge"
        self.kb_dir.mkdir()
        (self.kb_dir / "skin_care.md").write_text(
            "---\ncategory: 产品知识\nsource: 产品说明书\n---\n# 护肤指南\n\n## 刺痛与过敏处理\n使用产品后出现刺痛、发红、瘙痒等不适，应立即停用并清水冲洗，持续不适建议就医登记。\n\n## 保质期说明\n未开封产品保质期三年，开封后六到十二个月内使用完毕。",
            encoding="utf-8",
        )
        self.store = KnowledgeStore(self.dir / "kb.sqlite3", self.kb_dir)
        self.store.ensure_index()

    def tearDown(self):
        self._tmp.cleanup()

    def test_colloquial_query_reaches_knowledge_via_rewrite(self):
        """用户口语'脸疼'经改写命中知识库的'刺痛'章节；metadata 随结果返回。"""
        result = self.store.search("这个用了脸疼怎么办", top_k=2)
        self.assertTrue(result)
        self.assertIn("刺痛", result[0].text)
        self.assertEqual(result[0].metadata.get("category"), "产品知识")
        self.assertEqual(result[0].metadata.get("source"), "产品说明书")

    def test_hybrid_and_rerank_keep_exact_section_first(self):
        """混合通道（语义+标题+关键词覆盖）精排：标题与关键词双命中的章节排第一。"""
        chunks = self.store.search("保质期是多久", top_k=2)
        self.assertTrue(chunks)
        self.assertIn("保质期", chunks[0].section)
        self.assertIn("保质期", chunks[0].text)

    def test_ingest_excel_and_database_sources_into_same_store(self):
        """M17.1：Excel 与 Database 来源经统一管道进同一向量库并可检索。"""
        from openpyxl import Workbook

        xlsx = self.dir / "activity_rules.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["text", "category"])
        sheet.append(["618 预售付定金立减五十元。", "活动规则"])
        workbook.save(xlsx)
        workbook.close()
        self.store.ingest_file(xlsx)

        db_path = self.dir / "source.sqlite3"
        conn = sqlite3.connect(db_path)
        conn.execute("CREATE TABLE policy_rules (rule_text TEXT)")
        conn.execute("INSERT INTO policy_rules VALUES ('赠品随主订单一同发出，不单独寄送。')")
        conn.commit()
        conn.close()
        self.store.ingest_database(db_path, "policy_rules", "rule_text", category="售后规则")

        activity = self.store.search("预售有优惠吗", top_k=1)
        self.assertTrue(activity)
        self.assertEqual(activity[0].domain, "activity_rules")
        gift = self.store.search("赠品怎么发", top_k=1)
        self.assertTrue(gift)
        self.assertEqual(gift[0].domain, "policy_rules")

    def test_retriever_independent_of_store(self):
        """M18：KnowledgeRetriever 只依赖注入的候选函数，可脱离 KnowledgeStore 组装。"""
        from app.knowledge_base import KnowledgeRetriever

        def fake_candidates(query, domain, pool, threshold):
            return [{"chunk_id": "x:0", "domain": "x", "doc_title": "退款规则", "section": "到账时效", "content": "退款一般一到三个工作日到账。", "metadata": "{}", "base_score": 0.9}]

        retriever = KnowledgeRetriever(fake_candidates)
        result = retriever.retrieve("退款多久到账", top_k=1)
        self.assertEqual(result["chunks"][0].doc_title, "退款规则")
        self.assertEqual(result["chunks"][0].metadata, {})
        self.assertIn("rewrite", result)


class ColloquialAdverseEventRoutingTests(unittest.TestCase):
    """M18.1 口语不良反应表述与安全优先原则：'脸疼'类口语走业务工具而非 unknown/RAG。"""

    @classmethod
    def setUpClass(cls):
        from app.business_tools import BusinessTools

        root = Path(__file__).resolve().parents[1]
        cls.tools = BusinessTools(root / "data/beauty_cs_agent.sqlite3")

    def test_colloquial_face_pain_routes_to_adverse_reaction_tool(self):
        from app.action_registry import classify_copilot_intent, handle_copilot_chat

        self.assertEqual(classify_copilot_intent("这个用了脸疼怎么办"), "business_query")
        response = handle_copilot_chat(self.tools, "S00082", "这个用了脸疼怎么办")
        self.assertIn("get_adverse_reaction_record", response["tools_used"])
        self.assertNotIn("参考来源", response["message"])

    def test_colloquial_symptoms_classification(self):
        from app.action_registry import classify_copilot_intent

        self.assertEqual(classify_copilot_intent("用完脸疼得难受"), "business_query")
        self.assertEqual(classify_copilot_intent("脸上有点红肿"), "business_query")


if __name__ == "__main__":
    unittest.main()
