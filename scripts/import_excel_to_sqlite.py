#!/usr/bin/env python3
"""Deterministically import the competition workbook into SQLite.

All imported cell values are stored as TEXT. This keeps identifiers exact and
prevents Excel numeric IDs from being rendered in scientific notation or losing
precision. Dates and numeric values are normalized to stable text forms.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sqlite3
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl

TABLE_NAMES = {
    "数据说明": "data_description",
    "聊天记录": "chat_messages",
    "订单": "orders",
    "补发换货工单": "replacement_work_orders",
    "线下打款工单": "offline_payment_work_orders",
    "物流工单": "logistics_work_orders",
    "不良反应工单": "adverse_reaction_work_orders",
    "售后退货工单": "after_sales_return_work_orders",
}

HEADER_NAMES = {
    "会话ID": "conversation_id", "消息序号": "message_sequence", "发送时间": "sent_at",
    "角色": "role", "买家昵称": "buyer_nickname", "发送方": "sender", "店铺": "store",
    "内容类型": "content_type", "关联订单号": "order_number", "关联工单号": "work_order_number",
    "订单号": "order_number", "商品货号": "product_sku", "商品名称": "product_name",
    "数量": "quantity", "单价(元)": "unit_price_yuan", "实付金额(元)": "paid_amount_yuan",
    "订单状态": "order_status", "下单时间": "ordered_at", "付款时间": "paid_at", "发货时间": "shipped_at",
    "快递公司": "carrier", "物流单号": "tracking_number", "收货省": "shipping_province",
    "收货市": "shipping_city", "赠品": "gift", "买家留言": "buyer_note",
    "工单号": "work_order_number", "工单类型": "work_order_type", "售后原因": "after_sales_reason",
    "发出商品货号": "sent_product_sku", "发出商品名称": "sent_product_name",
    "原订单物流单号": "original_tracking_number", "补发物流单号": "replacement_tracking_number",
    "发货仓库": "fulfillment_warehouse", "客诉加急": "complaint_urgent", "工单状态": "work_order_status",
    "处理人": "handler", "创建时间": "created_at", "完成时间": "completed_at",
    "打款类型": "payment_type", "退款问题类型": "refund_issue_type", "退款金额(元)": "refund_amount_yuan",
    "支付宝实名": "alipay_real_name", "支付宝账号": "alipay_account", "相关物流单号": "related_tracking_number",
    "转账状态": "transfer_status", "问题类型": "issue_type", "问题包裹物流单号": "problem_tracking_number",
    "发货仓": "shipping_warehouse", "订单实付(元)": "order_paid_yuan", "处理方案": "resolution",
    "类型": "reaction_type", "年龄": "age", "肤质": "skin_type", "使用商品": "used_product",
    "产品批次号": "product_batch_number", "不适部位": "affected_area", "症状描述": "symptom_description",
    "用后多久出现": "onset_after_use", "是否停用": "stopped_use", "是否就医": "sought_medical_care",
    "任务状态": "task_status", "包裹类型": "package_type", "退货原因": "return_reason",
    "退货物流单号": "return_tracking_number", "退款编号": "refund_number", "签收建议": "receipt_recommendation",
    "是否异常": "is_abnormal", "message_id": "message_id", "scene_major": "scene_major",
    "scene_minor": "scene_minor", "is_target_buyer_message": "is_target_buyer_message",
    "message_text": "message_text", "chat_content": "chat_content", "category": "category", "image_path": "image_path",
}


def normalize_value(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        # Decimal(str(...)) avoids binary artefacts and fixed-point formatting
        # avoids scientific notation for long identifiers.
        return format(Decimal(str(value)), "f")
    return str(value)


def safe_column_names(headers: list[Any]) -> list[str]:
    result: list[str] = []
    seen: dict[str, int] = {}
    for index, header in enumerate(headers, 1):
        raw = normalize_value(header) or f"column_{index}"
        name = HEADER_NAMES.get(raw) or re.sub(r"[^0-9A-Za-z_]+", "_", raw).strip("_").lower() or f"column_{index}"
        if name[0].isdigit():
            name = f"column_{name}"
        count = seen.get(name, 0) + 1
        seen[name] = count
        result.append(name if count == 1 else f"{name}_{count}")
    return result


def quote(identifier: str) -> str:
    return '"' + identifier.replace('"', '""') + '"'


def workbook_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def import_workbook(xlsx_path: Path, db_path: Path) -> dict[str, Any]:
    xlsx_path = xlsx_path.resolve()
    db_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    source_hash = workbook_hash(xlsx_path)
    summary: dict[str, Any] = {"source": str(xlsx_path), "sha256": source_hash, "tables": {}}

    conn = sqlite3.connect(db_path)
    try:
        conn.execute("PRAGMA foreign_keys = ON")
        conn.execute("PRAGMA journal_mode = WAL")
        conn.execute("""CREATE TABLE IF NOT EXISTS import_runs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            source_path TEXT NOT NULL,
            source_sha256 TEXT NOT NULL,
            imported_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            table_summary_json TEXT NOT NULL
        )""")
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            if sheet_name == "数据说明":
                headers = [f"column_{i}" for i in range(1, sheet.max_column + 1)]
                data_rows = rows
            else:
                headers = safe_column_names(list(rows[0]))
                data_rows = rows[1:]
            table_name = TABLE_NAMES.get(sheet_name, re.sub(r"\W+", "_", sheet_name).lower())
            conn.execute(f"DROP TABLE IF EXISTS {quote(table_name)}")
            columns = ", ".join(f"{quote(col)} TEXT" for col in headers)
            conn.execute(f"CREATE TABLE {quote(table_name)} (_row_number INTEGER NOT NULL, {columns})")
            placeholders = ", ".join(["?"] * (len(headers) + 1))
            insert_sql = f"INSERT INTO {quote(table_name)} (_row_number, {', '.join(quote(c) for c in headers)}) VALUES ({placeholders})"
            count = 0
            for row_number, row in enumerate(data_rows, 1 if sheet_name == "数据说明" else 2):
                values = list(row[: len(headers)]) + [None] * max(0, len(headers) - len(row))
                if not any(value is not None and str(value) != "" for value in values):
                    continue
                conn.execute(insert_sql, [row_number] + [normalize_value(value) for value in values])
                count += 1
            conn.execute(f"CREATE INDEX {quote('idx_' + table_name + '_row')} ON {quote(table_name)} (_row_number)")
            summary["tables"][table_name] = {"sheet": sheet_name, "rows": count, "columns": headers}
        conn.execute(
            "INSERT INTO import_runs(source_path, source_sha256, table_summary_json) VALUES (?, ?, ?)",
            (str(xlsx_path), source_hash, json.dumps(summary["tables"], ensure_ascii=False, sort_keys=True)),
        )
        conn.commit()
    finally:
        conn.close()
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", type=Path, default=Path("data/raw/赛题 1：数据共情者-业务数据.xlsx"))
    parser.add_argument("--db", type=Path, default=Path("data/beauty_cs_agent.sqlite3"))
    args = parser.parse_args()
    print(json.dumps(import_workbook(args.xlsx, args.db), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
