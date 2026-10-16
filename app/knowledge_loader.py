"""M17.1 统一知识接入层（Knowledge Loader）。

多源知识统一转为 Document（page_content + metadata）后进入知识管道：

    文件 / 数据库 / API
      ↓ Loader
      ↓ Document(text + metadata)
      ↓ KnowledgeStore（chunk → embedding → vector store）

来源支持：
- Markdown：产品说明书 / 售后政策 / 客服规范（支持 YAML 风格 front-matter 声明 metadata）
- Excel：成分表 / 活动规则等表格知识（openpyxl，requirements 已有；表头行含 text 列，其余列进 metadata）
- Database：只读 SQL 提取政策配置类数据（不接入任何客户/订单数据）
- API：注入式 fetch 回调（无网络演示环境默认返回空；测试可注入 fake）

约束：知识库只存通用知识与政策，Loader 层不提供任何客户数据入口。
"""
from __future__ import annotations

import re
import sqlite3
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable


@dataclass(frozen=True)
class Document:
    """M17.2：知识文档 = page_content + metadata；回答可追溯"结论来自哪里"。"""

    text: str
    metadata: dict[str, Any] = field(default_factory=dict)

    def as_dict(self) -> dict[str, Any]:
        return {"text": self.text, "metadata": dict(self.metadata)}


_FRONT_MATTER = re.compile(r"\A---\s*\n(.*?)\n---\s*\n?", re.DOTALL)


def parse_front_matter(raw: str) -> tuple[dict[str, Any], str]:
    """解析 markdown 头部 front-matter（key: value 行），返回 (metadata, 剩余正文)。"""
    match = _FRONT_MATTER.match(raw)
    if not match:
        return {}, raw
    meta: dict[str, Any] = {}
    for line in match.group(1).splitlines():
        if ":" not in line:
            continue
        key, _, value = line.partition(":")
        key = key.strip()
        value = value.strip()
        if not key:
            continue
        if re.fullmatch(r"\d+", value):
            meta[key] = int(value)
        elif re.fullmatch(r"\d+\.\d+", value):
            meta[key] = float(value)
        else:
            meta[key] = value
    return meta, raw[match.end():]


class KnowledgeLoader:
    """统一 Loader：文件按扩展名分发，数据库/API 显式方法。"""

    def load_path(self, path: str | Path) -> list[Document]:
        path = Path(path)
        suffix = path.suffix.lower()
        if suffix == ".md":
            return self.load_markdown(path)
        if suffix in {".xlsx", ".xlsm"}:
            return self.load_excel(path)
        raise ValueError(f"unsupported knowledge source: {path.name}")

    def load_markdown(self, path: Path) -> list[Document]:
        """Markdown → 按 H2 切片；front-matter 提供文档级 metadata，切片继承。"""
        raw = path.read_text(encoding="utf-8")
        doc_meta, body = parse_front_matter(raw)
        title_match = re.search(r"^#\s+(.+)$", body, flags=re.MULTILINE)
        doc_title = title_match.group(1).strip() if title_match else path.stem
        documents: list[Document] = []
        parts = re.split(r"^##\s+(.+)$", body, flags=re.MULTILINE)
        preamble = parts[0].replace(doc_title, "", 1) if title_match else parts[0]
        preamble = re.sub(r"^#.*$|^>.*$", "", preamble, flags=re.MULTILINE).strip()
        if preamble:
            documents.append(Document(preamble, self._section_meta(doc_meta, path, doc_title, "概述")))
        for i in range(1, len(parts), 2):
            heading = parts[i].strip()
            section_body = re.sub(r"^>.*$", "", parts[i + 1], flags=re.MULTILINE).strip()
            if not section_body:
                continue
            for chunk_text in self._pack_section(section_body):
                documents.append(Document(chunk_text, self._section_meta(doc_meta, path, doc_title, heading)))
        return documents

    def load_excel(self, path: Path) -> list[Document]:
        """Excel → 每行一个 Document；表头含 text（或 content）列，其余列进 metadata。"""
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # requirements.txt 已声明，缺失时明确报错
            raise RuntimeError("openpyxl is required for Excel knowledge sources; run pip install -r requirements.txt") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        documents: list[Document] = []
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header: list[str] | None = None
            for row in rows:
                cells = ["" if cell is None else str(cell).strip() for cell in row]
                if header is None:
                    if not any(cells):
                        continue
                    header = [cell.lower() for cell in cells]
                    continue
                record = dict(zip(header, cells))
                text = record.get("text") or record.get("content") or ""
                if not text:
                    continue
                metadata = {key: value for key, value in record.items() if key not in {"text", "content"} and value}
                metadata.setdefault("source", path.name)
                metadata.setdefault("category", path.stem)
                metadata.setdefault("doc_title", f"{path.stem}·{sheet.title}")
                metadata.setdefault("section", f"第{len(documents) + 1}条")
                documents.append(Document(text, metadata))
        workbook.close()
        return documents

    def load_database(self, db_path: str | Path, table: str, text_field: str, metadata_fields: tuple[str, ...] = (), category: str | None = None, where: str = "") -> list[Document]:
        """Database → 只读 SELECT 提取知识行；禁止指向客户/订单数据表（调用方约束）。"""
        conn = sqlite3.connect(str(db_path))
        conn.row_factory = sqlite3.Row
        try:
            sql = f"SELECT {text_field}{', ' + ', '.join(metadata_fields) if metadata_fields else ''} FROM {table}"
            if where:
                sql += f" WHERE {where}"
            rows = conn.execute(sql).fetchall()
        finally:
            conn.close()
        documents: list[Document] = []
        for index, row in enumerate(rows):
            text = str(row[text_field] or "").strip()
            if not text:
                continue
            metadata = {field_name: row[field_name] for field_name in metadata_fields if row[field_name] not in (None, "")}
            metadata.setdefault("source", f"database:{table}")
            metadata.setdefault("category", category or table)
            metadata.setdefault("doc_title", table)
            metadata.setdefault("section", f"row-{index + 1}")
            documents.append(Document(text, metadata))
        return documents

    def load_api(self, name: str, fetch: Callable[[], list[dict[str, Any]]] | None = None) -> list[Document]:
        """API → 注入式 fetch 回调；未提供回调（无网络演示环境）返回空列表。"""
        if fetch is None:
            return []
        documents: list[Document] = []
        for index, record in enumerate(fetch()):
            text = str(record.get("text") or "").strip()
            if not text:
                continue
            metadata = {key: value for key, value in (record.get("metadata") or {}).items() if value}
            metadata.setdefault("source", f"api:{name}")
            metadata.setdefault("category", name)
            metadata.setdefault("doc_title", name)
            metadata.setdefault("section", f"item-{index + 1}")
            documents.append(Document(text, metadata))
        return documents

    @staticmethod
    def _section_meta(doc_meta: dict[str, Any], path: Path, doc_title: str, section: str) -> dict[str, Any]:
        metadata = dict(doc_meta)
        metadata.setdefault("source", path.name)
        metadata.setdefault("category", path.stem)
        metadata.setdefault("doc_title", doc_title)
        metadata["section"] = section
        return metadata

    @staticmethod
    def _pack_section(body: str, limit: int = 360) -> list[str]:
        if len(body) <= limit:
            return [body]
        packed: list[str] = []
        buffer = ""
        for paragraph in (p.strip() for p in re.split(r"\n\s*\n", body) if p.strip()):
            if buffer and len(buffer) + len(paragraph) > limit:
                packed.append(buffer.strip())
                buffer = paragraph
            else:
                buffer = f"{buffer}\n{paragraph}" if buffer else paragraph
        if buffer.strip():
            packed.append(buffer.strip())
        return packed
